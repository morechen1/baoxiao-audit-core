"""Generate local binary parser fixtures without external network access."""

from pathlib import Path

import fitz
from docx import Document


def main() -> None:
    fixtures = Path("tests/fixtures")
    fixtures.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Demo PDF page one")
    pdf.new_page().insert_text((72, 72), "Demo PDF page two")
    pdf.save(fixtures / "demo.pdf")
    pdf.close()

    document = Document()
    document.add_heading("Demo DOCX", level=0)
    document.add_heading("Section", level=1)
    document.add_paragraph("Offline parser fixture.")
    document.save(fixtures / "demo.docx")


if __name__ == "__main__":
    main()
