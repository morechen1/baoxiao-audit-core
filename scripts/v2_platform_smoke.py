"""Exercise V2 upload, batch, report, and short-circuit paths without Provider calls."""

from __future__ import annotations

import io
import os
from typing import Any, cast

import fitz
import httpx
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.engine import make_url


def require_safe_runtime() -> None:
    url = make_url(os.environ.get("DATABASE_URL", ""))
    if url.get_backend_name() != "postgresql" or url.database not in {
        "baoxiao_v2_development",
        "baoxiao_contest_final",
    }:
        raise SystemExit("V2 smoke requires an explicitly allowed PostgreSQL database.")
    if os.environ.get("SEMANTIC_PARSER_ENABLED", "").lower() not in {"false", "0"}:
        raise SystemExit("V2 smoke must explicitly disable Semantic Parser Provider calls.")


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "说明"
    table.cell(0, 1).text = "保险利益以合同为准"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text: str | None) -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    payload = cast(bytes, document.tobytes())
    document.close()
    return payload


def expect(response: httpx.Response, status: int = 200) -> dict[str, Any]:
    if response.status_code != status:
        raise AssertionError(f"HTTP {response.status_code}: {response.text}")
    return cast(dict[str, Any], response.json())


def main() -> None:
    require_safe_runtime()
    from app.main import app

    client = TestClient(app)
    high = expect(
        client.post(
            "/api/v2/screenings/upload",
            data={"material_type": "advertisement"},
            files={"file": ("高风险.txt", "本产品保证收益，资金绝对安全。", "text/plain")},
        )
    )
    if high["finding_count"] < 1:
        raise AssertionError("TXT high-risk upload produced no Finding")
    html_report = client.get(high["report_endpoints"]["html"])
    json_report = client.get(high["report_endpoints"]["json"])
    if html_report.status_code != 200 or "AI语义解析用于生成风险候选" not in html_report.text:
        raise AssertionError("HTML report export failed")
    report = expect(json_report)
    serialized = json_report.text
    if not report["findings"] or any(
        secret in serialized for secret in ("LLM_API_KEY", "prompt_snapshot", "hidden_reasoning")
    ):
        raise AssertionError("JSON audit report is incomplete or unsafe")

    low = expect(
        client.post(
            "/api/v2/screenings/upload",
            data={"material_type": "product_introduction"},
            files={
                "file": (
                    "合规说明.md",
                    "# 合同说明\n保险利益、责任免除及退保安排以正式合同为准。",
                    "text/markdown",
                )
            },
        )
    )
    low_detail = expect(client.get(low["report_endpoints"]["detail"]))
    semantic = low_detail["evidence_evaluation_summary"]["semantic_parser"]
    if low["finding_count"] != 0 or semantic["provider_calls"] != 0:
        raise AssertionError("Low-risk short circuit or Provider call guard failed")

    text_pdf = expect(
        client.post(
            "/api/v2/screenings/upload",
            data={"material_type": "other"},
            files={
                "file": ("text.pdf", pdf_bytes("Insurance contract information"), "application/pdf")
            },
        )
    )
    if text_pdf["source_type"] != "pdf":
        raise AssertionError("Text PDF ingestion failed")
    scanned = client.post(
        "/api/v2/screenings/upload",
        data={"material_type": "other"},
        files={"file": ("scan.pdf", pdf_bytes(None), "application/pdf")},
    )
    if (
        scanned.status_code != 409
        or scanned.json()["error"]["code"] != "upload_pdf_no_extractable_text"
    ):
        raise AssertionError("Scanned PDF was not explicitly rejected")

    batch = expect(
        client.post(
            "/api/v2/batches",
            data={"material_type": "advertisement"},
            files=[
                ("files", ("一.txt", "本产品保证收益。", "text/plain")),
                ("files", ("二.md", "保险责任以合同为准。", "text/markdown")),
                (
                    "files",
                    (
                        "三.docx",
                        docx_bytes("提前退保可无条件全额返还已交保费。"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
            ],
        )
    )
    if batch["total"] != 3 or batch["succeeded"] != 3 or batch["progress"] != 100:
        raise AssertionError("Three-file batch did not complete")
    if not all(item.get("report_endpoints", {}).get("html") for item in batch["items"]):
        raise AssertionError("Batch result drill-down endpoints missing")

    partial = expect(
        client.post(
            "/api/v2/batches",
            data={"material_type": "other"},
            files=[
                ("files", ("正常.txt", "保险利益以合同为准。", "text/plain")),
                ("files", ("扫描.pdf", pdf_bytes(None), "application/pdf")),
            ],
        )
    )
    if partial["status"] != "partial" or partial["succeeded"] != 1 or partial["failed"] != 1:
        raise AssertionError("Batch partial failure isolation failed")
    print(
        {
            "txt": "PASS",
            "md": "PASS",
            "docx": "PASS",
            "pdf_text": "PASS",
            "pdf_scan_rejected": "PASS",
            "batch_size": 3,
            "batch_partial": "PASS",
            "html_report": "PASS",
            "json_report": "PASS",
            "provider_calls": 0,
        }
    )


if __name__ == "__main__":
    main()
