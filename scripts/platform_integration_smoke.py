"""Exercise platform capabilities against PostgreSQL without Provider calls."""

from __future__ import annotations

import io
import os
import time
from typing import Any, cast

import fitz
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.engine import make_url


def require_safe_runtime() -> None:
    url = make_url(os.environ.get("DATABASE_URL", ""))
    if url.get_backend_name() != "postgresql" or url.database != "baoxiao_contest_final":
        raise SystemExit("platform smoke requires baoxiao_contest_final PostgreSQL")
    if os.environ.get("SEMANTIC_PARSER_ENABLED", "").lower() not in {"false", "0"}:
        raise SystemExit("platform smoke requires SEMANTIC_PARSER_ENABLED=false")
    if os.environ.get("SEMANTIC_SCREENING_ENABLED", "").lower() not in {"false", "0"}:
        raise SystemExit("platform smoke requires SEMANTIC_SCREENING_ENABLED=false")


def expect(response: Any, status: int = 200) -> dict[str, Any]:
    if response.status_code != status:
        raise AssertionError(f"HTTP {response.status_code}: {response.text}")
    return cast(dict[str, Any], response.json())


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "合同说明"
    table.cell(0, 1).text = "保险利益以合同为准"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text: str | None) -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    payload = cast(bytes, document.tobytes())
    document.close()
    return payload


def projection(detail: dict[str, Any]) -> list[tuple[Any, ...]]:
    return [
        (
            row["rule_id"],
            row["severity"],
            row["matched_text"],
            row["raw_start_offset"],
            row["raw_end_offset"],
        )
        for row in detail["findings"]
    ]


def wait_batch(client: TestClient, payload: dict[str, Any]) -> dict[str, Any]:
    for _ in range(100):
        if payload["status"] in {"completed", "partial"}:
            return payload
        time.sleep(0.1)
        payload = expect(client.get(f"/api/platform/batches/{payload['batch_id']}"))
    raise AssertionError("batch timeout")


def main() -> int:
    require_safe_runtime()
    from app.main import app

    client = TestClient(app)
    cases = (
        (
            "High",
            "监管推荐本产品，保证收益8%，本金绝对安全，今天投保即可领取限量礼品。",
            "advertisement",
        ),
        (
            "Boundary",
            "资金使用灵活，如有需要可随时退保没有损失。具体权益和现金价值请以合同约定为准。",
            "sales_script",
        ),
        (
            "Low",
            "本材料仅作产品信息说明，保险责任、等待期、责任免除及退保安排以正式保险合同为准。",
            "product_introduction",
        ),
    )
    compatibility: dict[str, str] = {}
    for label, text, material_type in cases:
        platform = expect(
            client.post(
                "/api/platform/screenings/text",
                json={
                    "title": label,
                    "material_type": material_type,
                    "raw_text": text,
                    "source_label": "platform_smoke",
                },
            )
        )
        platform_detail = expect(client.get(platform["report_endpoints"]["detail"]))
        v1 = expect(
            client.post(
                "/api/v1/screenings",
                json={
                    "title": label,
                    "material_type": material_type,
                    "raw_text": text,
                    "source_label": "platform_smoke",
                },
            )
        )
        v1_detail = expect(client.get(v1["report_endpoints"]["screening"]))
        if projection(platform_detail) != projection(v1_detail):
            raise AssertionError(f"{label} differs from direct V1 path")
        if platform_detail["runtime"]["parser_calls"] != 0:
            raise AssertionError("Provider call detected")
        compatibility[label] = "PASS"

    txt = expect(
        client.post(
            "/api/platform/screenings/upload",
            data={"material_type": "advertisement"},
            files={"file": ("高风险.txt", "本产品保证收益。", "text/plain")},
        )
    )
    md = expect(
        client.post(
            "/api/platform/screenings/upload",
            data={"material_type": "product_introduction"},
            files={"file": ("说明.md", "# 合同说明\n保险利益以合同为准。", "text/markdown")},
        )
    )
    docx = expect(
        client.post(
            "/api/platform/screenings/upload",
            data={"material_type": "sales_script"},
            files={
                "file": (
                    "话术.docx",
                    docx_bytes("提前退保可无条件全额返还已交保费。"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    )
    pdf = expect(
        client.post(
            "/api/platform/screenings/upload",
            data={"material_type": "other"},
            files={"file": ("text.pdf", pdf_bytes("Insurance contract"), "application/pdf")},
        )
    )
    scanned = client.post(
        "/api/platform/screenings/upload",
        data={"material_type": "other"},
        files={"file": ("scan.pdf", pdf_bytes(None), "application/pdf")},
    )
    if scanned.status_code != 409:
        raise AssertionError("scanned PDF did not fail closed")

    long_text = ("保险利益以合同约定为准。\n" * 8500) + "本产品保证收益。"
    long_result = expect(
        client.post(
            "/api/platform/screenings/text",
            json={
                "title": "长文档",
                "material_type": "advertisement",
                "raw_text": long_text,
                "source_label": "platform_smoke",
            },
        )
    )
    long_detail = expect(client.get(long_result["report_endpoints"]["detail"]))
    if long_detail["runtime"]["document_chunks"] < 2:
        raise AssertionError("long document did not use document orchestration")
    for finding in long_detail["findings"]:
        start = finding["raw_start_offset"]
        end = finding["raw_end_offset"]
        if long_text[start:end] != finding["matched_text"]:
            raise AssertionError("long document offset remapping failed")

    report_json = expect(client.get(txt["report_endpoints"]["json"]))
    report_html = client.get(txt["report_endpoints"]["html"])
    serialized = str(report_json)
    if report_html.status_code != 200 or not report_json["findings"]:
        raise AssertionError("report export failed")
    if any(value in serialized for value in ("LLM_API_KEY", "prompt_snapshot", "hidden_reasoning")):
        raise AssertionError("report contains forbidden fields")

    batch = wait_batch(
        client,
        expect(
            client.post(
                "/api/platform/batches",
                data={"material_type": "other"},
                files=[
                    ("files", ("一.txt", "本产品保证收益。", "text/plain")),
                    ("files", ("二.md", "保险利益以合同为准。", "text/markdown")),
                    (
                        "files",
                        (
                            "三.docx",
                            docx_bytes("提前退保可全额返还。"),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        ),
                    ),
                ],
            ),
            202,
        ),
    )
    if batch["succeeded"] != 3:
        raise AssertionError("three-file batch failed")
    partial = wait_batch(
        client,
        expect(
            client.post(
                "/api/platform/batches",
                data={"material_type": "other"},
                files=[
                    ("files", ("正常.txt", "保险利益以合同为准。", "text/plain")),
                    ("files", ("扫描.pdf", pdf_bytes(None), "application/pdf")),
                ],
            ),
            202,
        ),
    )
    if partial["succeeded"] != 1 or partial["failed"] != 1:
        raise AssertionError("batch partial failure isolation failed")

    print(
        {
            "v1_compatibility": compatibility,
            "txt": txt["status"],
            "md": md["status"],
            "docx": docx["status"],
            "pdf": pdf["status"],
            "scanned_pdf": "REJECTED",
            "long_document_chunks": long_detail["runtime"]["document_chunks"],
            "offset_remapping": "PASS",
            "batch": "PASS",
            "batch_partial": "PASS",
            "html_report": "PASS",
            "json_report": "PASS",
            "provider_calls": 0,
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
