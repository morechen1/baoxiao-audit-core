from __future__ import annotations

import io
import subprocess
from pathlib import Path
from typing import Any, cast

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.core.exceptions import DocumentIngestionError
from app.main import app
from app.services.platform.batch import BatchReviewService
from app.services.platform.cache import CachingSemanticClaimParser, ExactSemanticResultCache
from app.services.platform.export import PlatformAuditExportService
from app.services.platform.ingestion import MaterialIngestionService, MaterialInput
from app.services.platform.long_document import segment_long_document, spans_are_duplicates
from app.services.platform.reports import PlatformReportService
from app.services.platform.store import BatchItem, ChunkRun, PlatformScreeningResult, RuntimeStore
from app.services.screening.engine import DeterministicComplianceRuleEngine
from app.services.screening.rules import load_ruleset
from app.services.screening.segmenter import segment_marketing_text
from app.services.screening.semantic_parser import SemanticParserOutcome

ROOT = Path(__file__).parents[2]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "保证收益"
    document.add_paragraph("最后一段")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes(text: str | None) -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    payload = cast(bytes, document.tobytes())
    document.close()
    return payload


def test_txt_markdown_docx_and_pdf_ingestion() -> None:
    service = MaterialIngestionService(Settings(max_upload_bytes=1024 * 1024))
    txt = service.ingest(
        filename="材料.txt",
        content=b"\xef\xbb\xbfguaranteed return",
        material_type="advertisement",
        content_type="text/plain",
    )
    markdown = service.ingest(
        filename="说明.md",
        content="# 产品\n不保证收益".encode(),
        material_type="product_introduction",
        content_type="text/markdown",
    )
    docx = service.ingest(
        filename="材料.docx",
        content=_docx_bytes(),
        material_type="other",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    pdf = service.ingest(
        filename="材料.pdf",
        content=_pdf_bytes("Guaranteed return disclosure"),
        material_type="other",
        content_type="application/pdf",
    )
    assert txt.raw_text == "guaranteed return"
    assert markdown.raw_text == "# 产品\n不保证收益"
    assert docx.raw_text.splitlines() == ["第一段", "字段\t保证收益", "最后一段"]
    assert "Guaranteed return" in pdf.raw_text


def test_ingestion_guards_scanned_pdf_size_extension_mime_and_filename() -> None:
    service = MaterialIngestionService(Settings(max_upload_bytes=1024))
    with pytest.raises(DocumentIngestionError, match="upload_pdf_no_extractable_text"):
        service.ingest(
            filename="扫描.pdf",
            content=_pdf_bytes(None),
            material_type="other",
            content_type="application/pdf",
        )
    with pytest.raises(DocumentIngestionError, match="upload_extension_not_allowed"):
        service.ingest(
            filename="脚本.exe",
            content=b"unsafe",
            material_type="other",
            content_type="application/octet-stream",
        )
    with pytest.raises(DocumentIngestionError, match="upload_size_invalid"):
        service.ingest(
            filename="过大.txt",
            content=b"x" * 1025,
            material_type="other",
            content_type="text/plain",
        )
    with pytest.raises(DocumentIngestionError, match="upload_content_type_not_allowed"):
        service.ingest(
            filename="材料.pdf",
            content=_pdf_bytes("text"),
            material_type="other",
            content_type="text/html",
        )
    safe = service.ingest(
        filename="../../安全.txt",
        content="合规说明".encode(),
        material_type="other",
        content_type="text/plain",
    )
    assert safe.source_filename == "安全.txt"


def test_long_document_segmentation_offsets_overlap_and_limit() -> None:
    text = "第一段。\n" + "保证收益。" * 30 + "\n最后一段。"
    chunks = segment_long_document(text, max_chars=80, overlap=12, max_chunks=20)
    assert len(chunks) > 1
    assert all(text[item.start_offset : item.end_offset] == item.text for item in chunks)
    assert all(item.end_offset - item.start_offset <= 80 for item in chunks)
    assert all(
        chunks[index].start_offset < chunks[index - 1].end_offset
        for index in range(1, len(chunks))
    )
    with pytest.raises(DocumentIngestionError, match="long_document_chunk_limit_exceeded"):
        segment_long_document(text, max_chars=40, overlap=5, max_chunks=1)


def test_duplicate_suppression_is_same_rule_and_high_overlap_only() -> None:
    assert spans_are_duplicates("rule", 10, 20, "rule", 10, 20)
    assert spans_are_duplicates("rule", 10, 20, "rule", 11, 20)
    assert not spans_are_duplicates("rule", 10, 20, "other", 10, 20)
    assert not spans_are_duplicates("rule", 10, 20, "rule", 30, 40)


class _FakeParser:
    def __init__(self) -> None:
        self.calls = 0

    def supplement(self, **_: Any) -> SemanticParserOutcome:
        self.calls += 1
        return SemanticParserOutcome((), {"status": "completed", "provider_calls": 1})


def test_exact_cache_hit_same_result_and_text_or_model_invalidates() -> None:
    cache = ExactSemanticResultCache()
    fake = _FakeParser()
    settings = Settings(
        semantic_parser_enabled=True,
        semantic_parser_cache_enabled=True,
        llm_model="model-a",
    )
    parser = CachingSemanticClaimParser(
        ruleset=load_ruleset(),
        settings=settings,
        parser=cast(Any, fake),
        cache=cache,
    )
    arguments = {
        "raw_text": "保证收益",
        "material_sha256": "a" * 64,
        "segments": [],
        "deterministic": [],
    }
    first = parser.supplement(**arguments)
    second = parser.supplement(**arguments)
    assert first == second
    assert fake.calls == 1
    assert parser.last_cache_state == "hit"
    parser.supplement(**{**arguments, "raw_text": "不保证收益"})
    assert fake.calls == 2
    other_model = CachingSemanticClaimParser(
        ruleset=load_ruleset(),
        settings=settings.model_copy(update={"llm_model": "model-b"}),
        parser=cast(Any, fake),
        cache=cache,
    )
    other_model.supplement(**arguments)
    assert fake.calls == 3


class _FakeV1Reports:
    def institution_report(self, _session: Any, run_id: int) -> dict[str, Any]:
        start = 2 if run_id == 1 else 0
        return {
            "findings": [
                {
                    "finding_key": "F001",
                    "finding_sha256": str(run_id),
                    "rule_id": "guaranteed_return_or_principal",
                    "severity": "high",
                    "matched_text": "保证收益",
                    "raw_start_offset": start,
                    "raw_end_offset": start + 4,
                    "explanation": "风险说明",
                    "review_question": "是否承诺？",
                    "remediation_template": "删除承诺",
                    "evidence_status": "supported",
                    "evidence": [],
                }
            ]
        }

    def run_detail(self, _session: Any, _run_id: int) -> dict[str, Any]:
        return {"evidence_evaluation_summary": {"semantic_parser": {"provider_calls": 0}}}

    def consumer_notice(self, _session: Any, _run_id: int) -> dict[str, Any]:
        return {"risk_prompts": [], "questions_to_ask": [], "evidence_links": []}


def test_document_report_remaps_offsets_and_deduplicates_overlap() -> None:
    material = MaterialInput(
        title="长文",
        material_type="other",
        raw_text="xx保证收益yy",
        source_filename="long.txt",
        source_format="txt",
    )
    result = PlatformScreeningResult(
        result_id="result",
        material=material,
        chunks=(ChunkRun(0, 0, 8, 1), ChunkRun(1, 2, 10, 2)),
        runtime={
            "parser_calls": 0,
            "cache_hits": 0,
            "latency_ms": 1,
            "provider_fail_closed": False,
        },
        created_at="now",
    )
    reports = PlatformReportService()
    reports.v1 = cast(Any, _FakeV1Reports())
    detail = reports.detail(cast(Any, None), result)
    assert detail["finding_count"] == 1
    assert detail["findings"][0]["raw_start_offset"] == 2
    assert material.raw_text[2:6] == detail["findings"][0]["matched_text"]


def test_boundary_and_zero_finding_use_backend_overall_risk_summary(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw_text = "资金使用灵活，如有需要可随时退保没有损失。具体权益和现金价值请以合同约定为准。"
    candidates = DeterministicComplianceRuleEngine(load_ruleset()).run(
        raw_text=raw_text,
        material_sha256="a" * 64,
        segments=segment_marketing_text(raw_text, "a" * 64),
    )
    assert [(item.rule_id, item.severity) for item in candidates] == [
        ("surrender_or_cash_value_misstatement", "high")
    ]

    result = PlatformScreeningResult(
        result_id="boundary",
        material=MaterialInput("语境边界案例", "sales_script", raw_text, "粘贴文本", "text"),
        chunks=(ChunkRun(0, 0, len(raw_text), 1),),
        runtime={
            "parser_calls": 0,
            "cache_hits": 0,
            "latency_ms": 1,
            "provider_fail_closed": False,
        },
        created_at="now",
    )
    rows = [
        {
            "rule_id": candidates[0].rule_id,
            "severity": candidates[0].severity,
            "evidence_status": "supported",
            "evidence": [],
        }
    ]
    reports = PlatformReportService()
    monkeypatch.setattr(reports, "_findings", lambda _session, _result: rows)
    boundary_summary = reports.summary(cast(Any, None), result)
    assert boundary_summary["risk_level"] == "high"

    monkeypatch.setattr(reports, "_findings", lambda _session, _result: [])
    low_summary = reports.summary(cast(Any, None), result)
    assert low_summary["risk_level"] == "low"


def test_batch_and_exports_reuse_backend_risk_summary(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    material = MaterialInput("语境边界案例", "sales_script", "随时退保没有损失", "材料.txt", "txt")
    runtime = {
        "document_chunks": 1,
        "parser_calls": 0,
        "cache_hits": 0,
        "latency_ms": 1,
        "provider_fail_closed": False,
    }
    local_store = RuntimeStore()
    result = local_store.add_result(material, [ChunkRun(0, 0, len(material.raw_text), 1)], runtime)
    backend_summary = {
        "platform_result_id": result.result_id,
        "risk_level": "high",
        "finding_count": 1,
    }

    class _BackendSummaryReports:
        def summary(self, _session: Any, _result: Any) -> dict[str, Any]:
            return backend_summary

    record = local_store.create_batch([BatchItem(0, material, "材料.txt")])
    record.items[0].status = "success"
    record.items[0].result_id = result.result_id
    record.status = "completed"
    monkeypatch.setattr("app.services.platform.batch.RUNTIME_STORE", local_store)
    monkeypatch.setattr("app.services.platform.batch.PlatformReportService", _BackendSummaryReports)
    batch = BatchReviewService.payload(record)
    assert batch["items"][0]["risk_level"] == backend_summary["risk_level"]

    finding = {
        "finding_key": "F001",
        "rule_id": "surrender_or_cash_value_misstatement",
        "severity": "high",
        "matched_text": "随时退保没有损失",
        "raw_start_offset": 0,
        "raw_end_offset": 8,
        "explanation": "风险说明",
        "evidence": [],
    }

    class _ExportReports:
        def detail(self, _session: Any, _result: Any) -> dict[str, Any]:
            return {
                "screening_run_ids": [1],
                "status": "completed",
                "finding_count": 1,
                "runtime": runtime,
            }

        def institution(self, _session: Any, _result: Any) -> dict[str, Any]:
            return {"findings": [finding], "evidence_summary": {}}

        def consumer(self, _session: Any, _result: Any) -> dict[str, Any]:
            return {}

        def summary(self, _session: Any, _result: Any) -> dict[str, Any]:
            return backend_summary

    exporter = PlatformAuditExportService()
    exporter.reports = cast(Any, _ExportReports())
    monkeypatch.setattr(
        exporter,
        "_artifacts",
        lambda _session, _run_ids: {"institution": [], "consumer": []},
    )
    report = exporter.build(cast(Any, None), result)
    assert report["screening"]["risk_level"] == backend_summary["risk_level"]
    assert report["screening"]["risk_level_label_zh"] == "高风险"
    exported_html = exporter.html(report)
    assert "<b>风险：</b>高风险" in exported_html
    assert "<b>风险等级：</b>高风险" in exported_html


def test_batch_state_tracks_partial_failure_without_aborting(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    material = MaterialInput("材料", "other", "说明", "材料.txt", "txt")
    service = BatchReviewService(Settings(platform_batch_concurrency=2))
    local_store = RuntimeStore()
    record = local_store.create_batch(
        [BatchItem(0, material, "成功.txt"), BatchItem(1, material, "失败.txt")]
    )
    monkeypatch.setattr("app.services.platform.batch.RUNTIME_STORE", local_store)

    def fake_one(item: BatchItem) -> str:
        if "失败" in item.label:
            raise DocumentIngestionError("test_failure")
        return "result-id"

    monkeypatch.setattr(service, "_one", fake_one)
    service.process(record.batch_id)
    assert record.status == "partial"
    assert [item.status for item in record.items] == ["success", "failed"]


def test_html_report_escapes_material_and_exposes_no_secret() -> None:
    report: dict[str, Any] = {
        "safety_statement": "AI候选，系统校验",
        "material": {
            "title": "<script>x</script>",
            "source_format": "txt",
            "raw_text": "保证收益",
        },
        "screening": {
            "risk_level": "high",
            "risk_level_label_zh": "高风险",
            "finding_count": 1,
            "runtime": {
                "document_chunks": 1,
                "parser_calls": 0,
                "cache_hits": 0,
                "latency_ms": 1,
            },
        },
        "findings": [
            {
                "finding_key": "F001",
                "rule_id": "guaranteed_return_or_principal",
                "severity": "high",
                "matched_text": "保证收益",
                "raw_start_offset": 0,
                "raw_end_offset": 4,
                "explanation": "风险说明",
                "evidence": [],
            }
        ],
        "institution": {"artifacts": []},
        "consumer": {"artifacts": []},
    }
    exported = PlatformAuditExportService().html(report)
    assert "&lt;script&gt;" in exported
    assert "guaranteed_return_or_principal" in exported
    assert "<b>风险：</b>高风险" in exported
    assert "<b>风险等级：</b>高风险" in exported
    assert "LLM_API_KEY" not in exported
    assert "prompt_snapshot" not in exported


def test_invalid_upload_is_fail_closed_not_500() -> None:
    response = TestClient(app).post(
        "/api/platform/screenings/upload",
        data={"material_type": "other"},
        files={"file": ("scan.pdf", _pdf_bytes(None), "application/pdf")},
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "upload_pdf_no_extractable_text"


def test_v1_core_integrity_script_passes() -> None:
    completed = subprocess.run(
        ["python3", "scripts/verify_v1_core_integrity.py"],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    assert "V1 CORE INTEGRITY VERIFIED" in completed.stdout
