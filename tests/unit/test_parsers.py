from pathlib import Path

import fitz
from docx import Document

from app.core.config import Settings
from app.models.enums import DataType, ReviewStatus
from app.services.collection import FileCollector
from app.services.parsing import DocxParser, HtmlParser, ParsingService, PdfParser
from app.services.review import ReviewService


def test_pdf_parser_preserves_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    pdf = fitz.open()
    for text in ("第一页测试内容", "第二页测试内容"):
        page = pdf.new_page()
        page.insert_text((72, 72), text, fontname="china-s")
    pdf.save(path)
    pdf.close()

    result = PdfParser().parse(path)

    assert [page.page_number for page in result.pages] == [1, 2]
    assert len(result.pages) == 2
    assert result.metadata["requires_ocr"] is False


def test_docx_parser_preserves_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("演示产品条款", level=0)
    document.add_heading("责任范围", level=1)
    document.add_paragraph("仅用于测试。")
    document.save(path)

    result = DocxParser().parse(path)

    assert result.title == "演示产品条款"
    assert "责任范围" in result.plain_text
    assert result.metadata["headings"]


def test_html_parser_removes_navigation_and_scripts(sample_html: Path) -> None:
    result = HtmlParser().parse(sample_html)

    assert result.title == "演示规则"
    assert "仅用于测试" in result.plain_text
    assert "无关导航" not in result.plain_text
    assert "danger" not in result.plain_text


def scanned_pdf(path: Path) -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 32, 32), False)
    pixmap.clear_with(180)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    pdf.save(path)
    pdf.close()


def test_scanned_pdf_is_quarantined_for_ocr(session, tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    scanned_pdf(path)
    settings = Settings(database_url="sqlite://", data_dir=tmp_path / "data")
    collector = FileCollector(settings)
    document, _ = collector.persist(
        session,
        collector.collect(path),
        DataType.REGULATION.value,
    )

    parsed = ParsingService(settings).parse_document(session, document)

    assert parsed.metadata["requires_ocr"] is True
    assert parsed.warnings == ["suspected_scanned_pdf"]
    assert document.parse_status == "requires_ocr"
    assert document.raw_text is None
    assert document.final_review_status == ReviewStatus.COLLECTED.value
    assert document.chunks == []


def test_scanned_pdf_does_not_enter_review_batch(session, tmp_path: Path) -> None:
    path = tmp_path / "scanned-review.pdf"
    scanned_pdf(path)
    settings = Settings(database_url="sqlite://", data_dir=tmp_path / "data")
    collector = FileCollector(settings)
    document, _ = collector.persist(
        session,
        collector.collect(path),
        DataType.REGULATION.value,
    )
    ParsingService(settings).parse_document(session, document)

    batch = ReviewService(settings).export_batch(
        session,
        DataType.REGULATION.value,
        "jsonl",
    )

    assert batch.record_count == 0
