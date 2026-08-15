from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any, cast

import fitz
import pytest
from docx import Document
from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.api.routes import v2 as v2_routes
from app.core.config import Settings
from app.core.exceptions import DocumentIngestionError, ScreeningError
from app.services.audit_export import AuditReportExportService
from app.services.document_ingestion import IngestedMaterial, MarketingDocumentIngestionService

ROOT = Path(__file__).parents[2]


def test_txt_and_markdown_ingestion() -> None:
    service = MarketingDocumentIngestionService(Settings(max_upload_bytes=1024))
    txt = service.ingest("材料.txt", "保证收益".encode())
    markdown = service.ingest("说明.md", "# 产品\n不保证收益".encode())
    assert txt[0].raw_text == "保证收益"
    assert txt[0].source_type == "txt"
    assert markdown[0].raw_text.startswith("# 产品")
    assert markdown[0].source_type == "md"


def test_docx_ingestion_preserves_paragraph_and_table_order() -> None:
    document = Document()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "保证收益"
    document.add_paragraph("最后一段")
    stream = io.BytesIO()
    document.save(stream)

    material = MarketingDocumentIngestionService().ingest("材料.docx", stream.getvalue())[0]
    assert material.raw_text.splitlines() == ["第一段", "字段\t保证收益", "最后一段"]


def test_text_pdf_ingestion_and_scanned_pdf_rejection() -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Guaranteed return disclosure")
    text_pdf = document.tobytes()
    document.close()
    material = MarketingDocumentIngestionService().ingest("材料.pdf", text_pdf)[0]
    assert "Guaranteed return" in material.raw_text

    blank = fitz.open()
    blank.new_page()
    blank_pdf = blank.tobytes()
    blank.close()
    with pytest.raises(DocumentIngestionError, match="upload_pdf_no_extractable_text"):
        MarketingDocumentIngestionService().ingest("扫描件.pdf", blank_pdf)


def test_csv_single_text_column_and_file_guards() -> None:
    service = MarketingDocumentIngestionService(Settings(max_upload_bytes=1024, max_batch_items=3))
    materials = service.ingest("批量.csv", "文案\n保证收益\n不保证收益\n".encode())
    assert [item.raw_text for item in materials] == ["保证收益", "不保证收益"]
    with pytest.raises(DocumentIngestionError, match="upload_extension_not_allowed"):
        service.ingest("脚本.exe", b"unsafe")
    with pytest.raises(DocumentIngestionError, match="upload_size_invalid"):
        service.ingest("过大.txt", b"x" * 1025)
    with pytest.raises(DocumentIngestionError, match="upload_content_type_not_allowed"):
        service.validate_content_type("材料.pdf", "text/html")
    service.validate_content_type("材料.pdf", "application/pdf")


@pytest.mark.asyncio
async def test_batch_partial_failure_does_not_abort(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_screen(
        session: Session, material: IngestedMaterial, material_type: str
    ) -> dict[str, object]:
        del session, material_type
        if "失败" in material.title:
            raise ScreeningError("screening_failed_for_test")
        return {"material": material.source_name, "status": "completed", "screening_run_id": 1}

    monkeypatch.setattr(v2_routes, "_screen", fake_screen)
    files = [
        UploadFile(filename="成功.txt", file=io.BytesIO("正常说明".encode())),
        UploadFile(filename="失败.txt", file=io.BytesIO("保证收益".encode())),
    ]
    result = await v2_routes.batch_screening(
        files=files,
        material_type="other",
        session=cast(Session, None),
    )
    assert result["status"] == "partial"
    assert result["succeeded"] == 1
    assert result["failed"] == 1


def test_html_report_escapes_material_and_contains_no_secret_or_prompt() -> None:
    report: dict[str, Any] = {
        "safety_statement": "AI候选，系统校验",
        "material": {"title": "<script>x</script>", "type": "txt", "raw_text": "保证收益"},
        "screening": {"risk_rating": "high", "finding_count": 1},
        "findings": [
            {
                "finding_key": "F001",
                "rule_id": "guaranteed_return_or_principal",
                "severity": "high",
                "matched_text": "保证收益",
                "raw_start_offset": 0,
                "raw_end_offset": 4,
                "explanation": "风险说明",
                "evidence": [],
            }
        ],
        "institution": {"status": "completed", "validation_status": "passed"},
        "consumer": {"status": "completed", "validation_status": "passed"},
    }
    exported = AuditReportExportService().html(report)
    assert "&lt;script&gt;" in exported
    assert "guaranteed_return_or_principal" in exported
    assert "API_KEY" not in exported
    assert "prompt_snapshot" not in exported


def test_v2_development_dataset_is_dev_only_and_balanced() -> None:
    dataset = ROOT / "data/evaluations/v2_development_v1/v2_development_v1.jsonl"
    rows = [json.loads(line) for line in dataset.read_text().splitlines() if line]
    assert len(rows) == 120
    assert len({row["text"] for row in rows}) == 120
    assert all(row["split"] == "DEV_ONLY" for row in rows)
    assert sum(bool(row["expected_rules"]) for row in rows) == 72
    assert sum(not row["expected_rules"] for row in rows) == 48
