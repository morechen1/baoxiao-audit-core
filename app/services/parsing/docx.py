from pathlib import Path

from docx import Document

from app.services.parsing.base import DocumentParser, ParsedDocument, ParsedPage


class DocxParser(DocumentParser):
    def parse(self, path: Path) -> ParsedDocument:
        document = Document(str(path))
        paragraphs: list[str] = []
        headings: list[dict[str, str | int]] = []
        title = path.stem
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.startswith(("Title", "标题")) and title == path.stem:
                title = text
            if style.startswith(("Heading", "标题")):
                headings.append({"index": index, "text": text, "style": style})
            paragraphs.append(text)
        plain_text = "\n".join(paragraphs)
        return ParsedDocument(
            title=title,
            plain_text=plain_text,
            pages=[ParsedPage(1, plain_text)],
            metadata={"headings": headings, "paragraph_count": len(paragraphs)},
            warnings=["docx_page_numbers_unavailable"],
        )
